#!/usr/bin/env python3
"""Genera Presentacion.pptx a partir del contenido del deck de TaskFlow Scrum.

Uso:
    pip install python-pptx
    python3 generate_pptx.py

Crea 'Presentacion.pptx' junto a este script. No inventa métricas: las cifras
quedan como '[completar]' para llenarse con datos reales de Jira.
"""
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Pt, Inches
except ImportError:
    raise SystemExit(
        "Falta python-pptx. Instala con:  pip install python-pptx\n"
        "Alternativas sin Python: Marp (VS Code) o pandoc (ver README.md)."
    )

# (título, [líneas de contenido])
SLIDES = [
    ("TaskFlow Scrum — Autenticación segura",
     ["Equipo Scrum · Proyecto KAN", "Fecha: [completar]"]),
    ("El problema",
     ["La aplicación de tareas era anónima: cualquiera veía y editaba todo.",
      "Sin identidad de usuario → sin privacidad ni responsabilidad.",
      "Objetivo: que cada persona vea y gestione solo sus tareas, de forma segura."]),
    ("La decisión arquitectónica clave",
     ["Un único emisor de tokens (JWT): Supabase.",
      "Supabase Auth registra, autentica y emite el JWT.",
      "React gestiona la sesión y envía Authorization: Bearer <token>.",
      "Spring Boot actúa como Resource Server: valida, no emite.",
      "Evita el antipatrón de 'dos emisores de JWT'."]),
    ("Flujo de autenticación",
     ["Usuario → React → Supabase Auth (signUp / signInWithPassword)",
      "Supabase emite el JWT.",
      "React guarda la sesión y envía Authorization: Bearer <token>.",
      "Spring Boot valida firma (JWKS), issuer y expiración.",
      "Extrae identidad del claim 'sub' y filtra tareas por user_id."]),
    ("Seguridad por diseño",
     ["El backend nunca confía en un userId del frontend → usa solo el claim 'sub'.",
      "Aislamiento horizontal: findByIdAndUserId + 404 para recursos ajenos.",
      "Credenciales fuera del código (variables de entorno) — KAN-15.",
      "Sin service_role_key en frontend · sin tokens en logs · sin stacktraces al cliente."]),
    ("Alcance del trabajo: 14 tareas",
     ["Base backend: KAN-15, KAN-14",
      "Decisión de auth: KAN-16, KAN-17, KAN-18",
      "Frontend: KAN-29, KAN-30, KAN-31, KAN-32, KAN-33",
      "Integración segura: KAN-34, KAN-35, KAN-36, KAN-37"]),
    ("Cómo trabaja el equipo",
     ["Cada tarea = paquete reproducible de 9 documentos.",
      "Convenciones Git: rama por tarea, Conventional Commits, PR por Jira key.",
      "Frontend y backend nunca comparten rama.",
      "Guías de Git separadas para Windows (sin &&) y macOS/Linux."]),
    ("Orden de implementación",
     ["KAN-15 → KAN-14 → KAN-29 → KAN-30/31 → KAN-32 → KAN-33",
      "KAN-34 → KAN-35 → KAN-36 → KAN-37",
      "Dependencias mapeadas para evitar bloqueos entre desarrolladores."]),
    ("Estado y métricas",
     ["(Rellenar con datos reales de Jira antes de presentar)",
      "Sprints ejecutados: [completar]",
      "Velocidad promedio: [completar] SP",
      "% cumplimiento del último sprint: [completar]",
      "Tareas cerradas / totales: [completar]"]),
    ("Riesgos gestionados",
     ["JDK 17 en máquina vs. Java 21 requerido → [estado]",
      "backend/target/ versionado en Git → [estado]",
      "Credenciales en texto plano → resuelto (KAN-15) [confirmar]",
      "Proyecto Supabase aún no creado → prerrequisito para pruebas."]),
    ("Próximos pasos",
     ["1. Crear el proyecto Supabase (URL + JWKS + claves).",
      "2. Ejecutar los paquetes en orden (EJECUTAR KAN-XX).",
      "3. Verificar pruebas de aislamiento con dos usuarios.",
      "4. Capturar métricas reales y evidencias."]),
    ("Gracias",
     ["TaskFlow Scrum — Autenticación segura",
      "Repositorio: MalcomJesid/taskflow-scrum · Jira: proyecto KAN"]),
]


def build():
    prs = Presentation()
    # Slide 0: portada con layout de título
    title_layout = prs.slide_layouts[0]
    content_layout = prs.slide_layouts[1]

    first = SLIDES[0]
    s = prs.slides.add_slide(title_layout)
    s.shapes.title.text = first[0]
    s.placeholders[1].text = "\n".join(first[1])

    for title, bullets in SLIDES[1:]:
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.clear()
        for i, line in enumerate(bullets):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = line
            p.font.size = Pt(18)

    out = Path(__file__).with_name("Presentacion.pptx")
    prs.save(out)
    print(f"Generado: {out}")


if __name__ == "__main__":
    build()
